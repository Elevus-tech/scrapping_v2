from openpyxl import load_workbook


def carregar_produtos(path):
    """
    Carrega a planilha e retorna:
    - workbook
    - worksheet
    - lista de produtos

    Estrutura esperada:
    A = ID WooCommerce
    B = SKU
    C = Nome Normalizado (ignorado)
    D = Nome Atual (usado na busca)
    E = Preço
    """

    wb = load_workbook(path)

    ws = wb.active

    produtos = []

    for row in range(2, ws.max_row + 1):

        woo_id = ws[f"A{row}"].value
        sku = ws[f"B{row}"].value
        nome = ws[f"D{row}"].value

        if not nome:
            continue

        produtos.append({
            "row": row,
            "id": woo_id,
            "sku": str(sku).strip() if sku else "",
            "nome": str(nome).strip()
        })

    return wb, ws, produtos


def salvar_preco(ws, row, preco):
    """
    Salva o preço encontrado na coluna E.
    """

    ws[f"E{row}"] = preco